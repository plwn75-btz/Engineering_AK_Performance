"""
MDR Workflow Data Processor
Integrates:
1. MM-ASK-1A-GEN01-ENG-MDR-0001_B1- Cut off 18-Sep-26.xlsx (5 tabs)
2. ASK Technical Document Log Report - 18 Sep 26.xlsx
Generates: mdr_data.js for index_mdr.html
"""

import os
import re
import sys
import glob
import json
from pathlib import Path
from datetime import datetime, timedelta, date
import openpyxl
import pandas as pd

def calc_working_days(d1, d2):
    if d1 is None or d2 is None: return None
    if d2 < d1: return 0
    cur = d1 + timedelta(days=1)
    work_days = 0
    while cur <= d2:
        if cur.weekday() < 5:  # Monday to Friday (0-4)
            work_days += 1
        cur += timedelta(days=1)
    return work_days

def process_mdr(mdr_path=None, log_path=None):
    if not mdr_path:
        files = sorted([f for f in glob.glob("*MDR*.xlsx") if not os.path.basename(f).startswith("~$")], reverse=True)
        if not files:
            print("Error: No MDR Excel file found.")
            return False
        mdr_path = files[0]

    if not log_path:
        files = sorted([f for f in glob.glob("ASK Technical Document Log Report*.xlsx") if not os.path.basename(f).startswith("~$")], reverse=True)
        if not files:
            print("Error: No Technical Document Log Excel file found.")
            return False
        log_path = files[0]

    print(f"Loading Log Report: {log_path}...")
    wb_log = openpyxl.load_workbook(log_path, data_only=True)
    sheet_log = wb_log['Page 1']
    log_vals = sheet_log.values
    for _ in range(9): next(log_vals)
    headers_log = next(log_vals)
    df_log = pd.DataFrame(log_vals, columns=headers_log)
    j_log = df_log[df_log['Submitted By'].astype(str).str.contains('zm167dc@jutal.com', case=False, na=False)].copy()

    doc_log_map = {}
    for _, row in j_log.iterrows():
        d_no = str(row['Document No.'] or '').strip()
        if not d_no: continue
        if d_no not in doc_log_map:
            doc_log_map[d_no] = []
        
        in_str = str(row['Actual Submit In'] or '').strip()
        out_str = str(row['Actual Submit Out'] or '').strip()
        
        in_dt = None
        out_dt = None
        dur = None
        if in_str:
            try: in_dt = datetime.strptime(in_str, '%d-%b-%Y %H:%M:%S')
            except: pass
        if out_str:
            try: out_dt = datetime.strptime(out_str, '%d-%b-%Y %H:%M:%S')
            except: pass
        if in_dt and out_dt:
            dur = calc_working_days(in_dt.date(), out_dt.date())

        doc_log_map[d_no].append({
            'rev': str(row['Rev.'] or '').strip(),
            'submittal_id': str(row['Submittal ID'] or '').strip(),
            'doc_category': str(row.get('Document Category') or '').strip(),
            'submit_in_raw': in_str,
            'submit_out_raw': out_str,
            'in_dt': in_dt.strftime('%Y-%m-%d') if in_dt else None,
            'out_dt': out_dt.strftime('%Y-%m-%d') if out_dt else None,
            'duration_days': dur,
            'status_code': str(row['Return Status Code'] or '').strip(),
            'status_name': str(row['Return Status Name'] or '').strip()
        })

    print(f"Indexed {len(doc_log_map)} documents from log.")

    print(f"Loading MDR: {mdr_path}...")
    wb_mdr = openpyxl.load_workbook(mdr_path, data_only=True)
    sheets = ['WP01 TOPSIDE', 'WP01 JACKET', 'WP02 TOPSIDE', 'WP02 JACKET', 'PRO.ENGINEERING_MR TBE']

    def parse_date_val(val):
        if val is None: return None
        if isinstance(val, datetime):
            return val.strftime('%Y-%m-%d')
        s = str(val).strip()
        if s.upper() in ['N/A', '-', '']:
            return 'N/A'
        for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d-%b-%Y']:
            try:
                return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
            except:
                pass
        return s

    records = []
    for sname in sheets:
        sheet = wb_mdr[sname]
        header_row = 11 if sname == 'PRO.ENGINEERING_MR TBE' else 13
        
        for r in range(header_row + 1, sheet.max_row + 1):
            doc_no_val = sheet.cell(row=r, column=9).value
            if not doc_no_val: continue
            doc_no = str(doc_no_val).strip()
            if not doc_no.startswith('MM-'): continue
            
            title = str(sheet.cell(row=r, column=10).value or '').strip()
            disc = str(sheet.cell(row=r, column=4).value or '').strip()
            area = str(sheet.cell(row=r, column=3).value or '').strip()
            doc_type = str(sheet.cell(row=r, column=5).value or '').strip()
            doc_class = str(sheet.cell(row=r, column=11).value or '').strip()
            
            ifr_plan = parse_date_val(sheet.cell(row=r, column=23).value)
            ifa_plan = parse_date_val(sheet.cell(row=r, column=29).value)
            afc_plan = parse_date_val(sheet.cell(row=r, column=35).value)
            ap_plan = parse_date_val(sheet.cell(row=r, column=40).value)
            ap_code = str(sheet.cell(row=r, column=42).value or '').strip()
            
            logs = doc_log_map.get(doc_no, [])
            logs.sort(key=lambda x: x['in_dt'] or '9999-99-99')
            
            ifr_log = next((l for l in logs if l['rev'].startswith('A')), None)
            ifa_log = next((l for l in logs if l['rev'].startswith('B')), None)
            afc_log = next((l for l in logs if l['rev'].startswith('C')), None)
            
            if not ifr_log and len(logs) > 0: ifr_log = logs[0]
            if not ifa_log and ifa_plan != 'N/A' and len(logs) > 1: ifa_log = logs[1]
            if not afc_log and len(logs) > 2: afc_log = logs[2]
            
            # Contractor incorporation turnaround (Working Days: Monday to Friday)
            ifa_incorp_days = None
            if ifa_log and ifa_log['in_dt'] and ifr_log and ifr_log['out_dt']:
                try:
                    dt_in = datetime.strptime(ifa_log['in_dt'], '%Y-%m-%d')
                    dt_out = datetime.strptime(ifr_log['out_dt'], '%Y-%m-%d')
                    ifa_incorp_days = calc_working_days(dt_out.date(), dt_in.date())
                except: pass
                
            afc_incorp_days = None
            prev_out = ifa_log['out_dt'] if (ifa_log and ifa_log['out_dt']) else (ifr_log['out_dt'] if ifr_log else None)
            if afc_log and afc_log['in_dt'] and prev_out:
                try:
                    dt_in = datetime.strptime(afc_log['in_dt'], '%Y-%m-%d')
                    dt_out = datetime.strptime(prev_out, '%Y-%m-%d')
                    afc_incorp_days = calc_working_days(dt_out.date(), dt_in.date())
                except: pass
                
            # AFC -> AP incorporation turnaround (Working Days: Monday to Friday)
            ap_incorp_days = None
            ap_out_dt = None
            if afc_log and afc_log['out_dt']:
                after_afc_logs = [l for l in logs if l['in_dt'] and l['in_dt'] > afc_log['out_dt']]
                if after_afc_logs:
                    try:
                        dt_in = datetime.strptime(after_afc_logs[0]['in_dt'], '%Y-%m-%d')
                        dt_out = datetime.strptime(afc_log['out_dt'], '%Y-%m-%d')
                        diff = calc_working_days(dt_out.date(), dt_in.date())
                        if diff is not None and diff >= 0:
                            ap_incorp_days = diff
                    except: pass
                    appr_log = next((l for l in after_afc_logs if 'APPR' in (l['status_code'] or '').upper()), None)
                    if appr_log:
                        ap_out_dt = appr_log['out_dt']
                elif afc_log['status_code'] and 'APPR' in afc_log['status_code'].upper():
                    ap_incorp_days = 0
                    ap_out_dt = afc_log['out_dt']

            # Determine Document Category (Part B - Engineering Doc vs Part C - Engineering Dwg)
            cat = None
            for l in logs:
                if l.get('doc_category'):
                    cat = l['doc_category']
                    break
            if not cat:
                if any(x in doc_no for x in ['-DWG-', '-PID-', '-PFD-', '-UFD-', '-LAY-', '-SCH-', '-SLD-', '-ISO-']):
                    cat = 'Part C - Engineering Dwg'
                elif doc_type in ['DWG', 'PID', 'PFD', 'UFD', 'LAY', 'SCH', 'SLD', 'ISO', 'GA']:
                    cat = 'Part C - Engineering Dwg'
                else:
                    cat = 'Part B - Engineering Doc'

            clean_category = 'Part C - Engineering Dwg' if 'Part C' in cat else 'Part B - Engineering Doc'
            is_approved = (ap_code.upper() in ['APPR', 'AP', 'CODE 1', 'CODE 2']) or any('APPR' in (l['status_code'].upper()) for l in logs)

            records.append({
                'doc_no': doc_no,
                'title': title,
                'work_package': sname,
                'discipline': disc,
                'area': area,
                'doc_type': doc_type,
                'doc_class': doc_class,
                'doc_category': clean_category,
                
                'ifr': {
                    'plan': ifr_plan,
                    'rev': ifr_log['rev'] if ifr_log else None,
                    'submit_in': ifr_log['in_dt'] if ifr_log else None,
                    'submit_out': ifr_log['out_dt'] if ifr_log else None,
                    'review_days': ifr_log['duration_days'] if ifr_log else None,
                    'status_code': ifr_log['status_code'] if ifr_log else None,
                    'status_name': ifr_log['status_name'] if ifr_log else None,
                    'submittal_id': ifr_log['submittal_id'] if ifr_log else None
                },
                
                'ifa': {
                    'is_skipped': (ifa_plan == 'N/A'),
                    'plan': ifa_plan,
                    'rev': ifa_log['rev'] if ifa_log else None,
                    'submit_in': ifa_log['in_dt'] if ifa_log else None,
                    'submit_out': ifa_log['out_dt'] if ifa_log else None,
                    'review_days': ifa_log['duration_days'] if ifa_log else None,
                    'incorp_days': ifa_incorp_days if (ifa_incorp_days is not None and ifa_incorp_days >= 0) else None,
                    'status_code': ifa_log['status_code'] if ifa_log else None,
                    'status_name': ifa_log['status_name'] if ifa_log else None,
                    'submittal_id': ifa_log['submittal_id'] if ifa_log else None
                },
                
                'afc': {
                    'plan': afc_plan,
                    'rev': afc_log['rev'] if afc_log else None,
                    'submit_in': afc_log['in_dt'] if afc_log else None,
                    'submit_out': afc_log['out_dt'] if afc_log else None,
                    'review_days': afc_log['duration_days'] if afc_log else None,
                    'incorp_days': afc_incorp_days if (afc_incorp_days is not None and afc_incorp_days >= 0) else None,
                    'status_code': afc_log['status_code'] if afc_log else None,
                    'status_name': afc_log['status_name'] if afc_log else None,
                    'submittal_id': afc_log['submittal_id'] if afc_log else None
                },
                
                'ap': {
                    'plan': ap_plan,
                    'is_approved': is_approved,
                    'code': ap_code,
                    'incorp_days': ap_incorp_days,
                    'submit_out': ap_out_dt,
                    'latest_status': logs[-1]['status_code'] if logs else None
                },
                
                'total_revisions_logged': len(logs)
            })

    # Extract cut-off date from MDR filename
    base_name = Path(mdr_path).name
    m = re.search(r'(\d{1,2})[-\s]([A-Za-z]{3})[-\s](\d{2,4})', base_name)
    if m:
        day = m.group(1).zfill(2)
        mon = m.group(2).capitalize()
        yr = m.group(3)
        if len(yr) == 2: yr = "20" + yr
        report_date = f"{day}-{mon}-{yr}"
    else:
        report_date = "18-Sep-2026"

    out_file = Path('mdr_data.js')
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(f'window.MDR_REPORT_DATE = "{report_date}";\n')
        f.write('window.MDR_DATA = ' + json.dumps(records, ensure_ascii=False) + ';')

    print(f"Successfully generated {out_file} with {len(records)} records (Cut-off Date: {report_date}).")
    return True

if __name__ == '__main__':
    process_mdr()
